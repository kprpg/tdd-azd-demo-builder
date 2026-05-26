from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTFILE = Path(__file__).resolve().parents[1] / "azure-arc-dba-briefing-deck.pptx"


TITLE_COLOR = RGBColor(11, 61, 145)
ACCENT_COLOR = RGBColor(0, 120, 212)
TEXT_COLOR = RGBColor(32, 32, 32)
MUTED_COLOR = RGBColor(90, 90, 90)
BG_COLOR = RGBColor(248, 250, 252)


def add_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.33), Inches(0.28)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = TITLE_COLOR
    band.line.fill.background()


def add_title(slide, title, subtitle=None):
    add_background(slide)
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.55), Inches(11.8), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Aptos Display"
    run.font.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = TITLE_COLOR

    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.68), Inches(1.25), Inches(11.2), Inches(0.5))
        sub_tf = sub_box.text_frame
        p = sub_tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        run.font.name = "Aptos"
        run.font.size = Pt(13)
        run.font.color.rgb = MUTED_COLOR


def add_bullets(slide, items, left=0.78, top=1.75, width=11.2, height=4.8, font_size=20):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT_COLOR
        p.space_after = Pt(8)


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(7.0), Inches(12), Inches(0.28))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED_COLOR


def add_two_column_slide(slide, left_title, left_items, right_title, right_items):
    left_panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.6), Inches(5.9), Inches(4.9))
    left_panel.fill.solid()
    left_panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    left_panel.line.color.rgb = RGBColor(220, 227, 235)

    right_panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.78), Inches(1.6), Inches(5.85), Inches(4.9))
    right_panel.fill.solid()
    right_panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    right_panel.line.color.rgb = RGBColor(220, 227, 235)

    left_head = slide.shapes.add_textbox(Inches(0.95), Inches(1.82), Inches(5.0), Inches(0.4))
    p = left_head.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = left_title
    r.font.name = "Aptos Display"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = ACCENT_COLOR

    right_head = slide.shapes.add_textbox(Inches(7.03), Inches(1.82), Inches(5.0), Inches(0.4))
    p = right_head.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = right_title
    r.font.name = "Aptos Display"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = ACCENT_COLOR

    add_bullets(slide, left_items, left=0.95, top=2.25, width=5.15, height=3.9, font_size=16)
    add_bullets(slide, right_items, left=7.03, top=2.25, width=5.0, height=3.9, font_size=16)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Azure Arc For SQL Server", "DBA pilot, SQL Server 2016 ESUs, demo value, and The Factory")
    add_bullets(
        slide,
        [
            "Hybrid SQL management control plane for on-premises and multicloud servers",
            "Focus: pilot onboarding, SQL 2016 support strategy, free vs paid capabilities",
            "Outcome: move from one-off onboarding to a governed operating model",
        ],
        top=2.15,
        height=2.8,
        font_size=22,
    )
    add_footer(slide, "Prepared for DBA discussion | May 2026")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Why DBAs Should Care")
    add_bullets(
        slide,
        [
            "Single Azure control plane for SQL estate outside Azure",
            "Central inventory of servers, instances, and databases",
            "Better governance, lifecycle visibility, and modernization planning",
            "ESU readiness becomes an operational workflow instead of a last-minute scramble",
        ],
    )
    add_footer(slide, "Azure Arc-enabled servers + SQL Server enabled by Azure Arc")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Pilot Scope And Prerequisites")
    add_two_column_slide(
        slide,
        "Pilot scope",
        [
            "5 to 15 representative database servers",
            "Include SQL Server 2016 and newer versions",
            "Mix of easy candidates and production-like systems",
            "Define clear success criteria before onboarding",
        ],
        "Prerequisites",
        [
            "Azure subscription and resource group",
            "Outbound HTTPS on TCP 443 or approved proxy",
            "Local admin rights for Connected Machine agent",
            "RBAC, tags, naming, and ownership model",
            "Enable SQL Server in Azure Arc after machine onboarding",
        ],
    )
    add_footer(slide, "Keep the pilot small enough to learn, broad enough to prove repeatability")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Pilot Workflow")
    add_bullets(
        slide,
        [
            "1. Select pilot servers and confirm ownership",
            "2. Connect hosts to Azure Arc as Arc-enabled servers",
            "3. Enable SQL Server enabled by Azure Arc",
            "4. Validate inventory, tags, RBAC, and health",
            "5. Decide whether the onboarding motion is ready to scale",
        ],
        font_size=21,
    )
    add_footer(slide, "Success is operational repeatability, not just technical connectivity")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "What SQL Server 2016 ESUs Require")
    add_bullets(
        slide,
        [
            "Arc-enabled server onboarding is the prerequisite",
            "SQL Server enabled by Azure Arc must be configured and healthy",
            "Choose the right ESU model: v-core, p-core, or p-core with unlimited virtualization",
            "Use active Software Assurance, SQL subscription, or Azure pay-as-you-go billing",
            "Classify each SQL 2016 workload: upgrade, migrate, retire, or ESU",
        ],
        font_size=20,
    )
    add_footer(slide, "ESUs are a temporary exception path, not the default lifecycle strategy")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Free Value For DBAs")
    add_bullets(
        slide,
        [
            "Central inventory of Arc-enabled servers and SQL instances",
            "Version, edition, core count, host OS, and database visibility",
            "Azure Resource Graph queries and dashboards for SQL estate reporting",
            "Governance through tags, RBAC, and policy attachment points",
            "Migration readiness and custom data collection through Arc Run Command",
        ],
        font_size=19,
    )
    add_footer(slide, "Lead the demo with immediate operational value before discussing paid services")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Paid And Add-On Value")
    add_two_column_slide(
        slide,
        "Paid capabilities",
        [
            "SQL Server ESU subscriptions",
            "Defender for Cloud / Defender for SQL scenarios",
            "SQL Server pay-as-you-go licensing",
            "Monitoring scenarios backed by AMA and Log Analytics",
        ],
        "Positioning guidance",
        [
            "Keep free and paid stories separate",
            "Reserve ESUs for justified exception workloads",
            "Use cost and lifecycle data to decide when to transition off ESUs",
            "Avoid presenting ESU as the default answer for all SQL 2016 systems",
        ],
    )
    add_footer(slide, "Paid options should sit on top of a well-governed Arc foundation")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Suggested Live Demo Flow")
    add_bullets(
        slide,
        [
            "Show Arc-enabled servers and highlight machine health, tags, and ownership",
            "Show SQL Server enabled by Azure Arc and compare a SQL 2016 instance with a newer version",
            "Show free value first: inventory, Resource Graph, dashboards, and governance",
            "Explain where ESU configuration fits without turning the demo into a licensing discussion",
            "Close with how The Factory makes this repeatable",
        ],
        font_size=19,
    )
    add_footer(slide, "Aim for a 15-20 minute demo with clear DBA relevance")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "How The Factory Helps")
    add_bullets(
        slide,
        [
            "Standardize discovery of the SQL estate",
            "Classify SQL 2016 workloads into upgrade, migrate, retire, or ESU",
            "Automate Arc onboarding and SQL extension enablement",
            "Apply consistent tags, RBAC, and policy",
            "Roll out in waves instead of one server at a time",
        ],
        font_size=20,
    )
    add_footer(slide, "The Factory turns a pilot into an operating model")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Recommended Next Steps")
    add_bullets(
        slide,
        [
            "Confirm the pilot server list and ownership model",
            "Validate network and proxy requirements",
            "Onboard the pilot and verify SQL visibility",
            "Create the SQL 2016 classification list",
            "Define the Factory-led rollout pattern for scale",
        ],
        font_size=21,
    )
    add_footer(slide, "Use the pilot to prove repeatability, governance, and DBA value")

    prs.save(OUTFILE)
    print(f"Created {OUTFILE}")


if __name__ == "__main__":
    build_deck()